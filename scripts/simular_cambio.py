"""Simula que llega el dato REAL de un mes en el Excel de cumplimiento (para demos y pruebas).

Escribe el real del mes indicado (y actualiza el acumulado de ese mes para que sea
consistente). Antes de modificar crea una copia de respaldo <archivo>.bak.

Ejemplo (desde la raíz del proyecto):
    python scripts/simular_cambio.py datos/entrada/Libro1.xlsx --segmento Pospago --mes 11 --real 800
"""
from __future__ import annotations

import argparse
import shutil
import sys
from pathlib import Path

RAIZ_PROYECTO = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(RAIZ_PROYECTO), str(RAIZ_PROYECTO / "nodos_hijo" / "nodo_excel_cumplimiento")]

import openpyxl  # noqa: E402

from comun.utilidades import resolver_ruta  # noqa: E402
from extraccion import (ETIQUETAS_POR_DEFECTO, Bloque, ParColumnas, detectar_bloques,  # noqa: E402
                        detectar_columnas, leer_celdas, numero_celda)


def escribir_real(hoja, celdas, bloque: Bloque, par: ParColumnas, real: float) -> None:
    f = bloque.filas
    meta_jugada = numero_celda(celdas, f["meta_jugada"], par.col_meta)
    linea_base = numero_celda(celdas, f["linea_base"], par.col_meta)
    hoja.cell(f["meta_driver"], par.col_real).value = real
    hoja.cell(f["meta_jugada"], par.col_real).value = real / meta_jugada
    hoja.cell(f["unidades_adicionales"], par.col_real).value = real - linea_base


def main() -> int:
    parser = argparse.ArgumentParser(description="Escribe el real de un mes en el Excel de cumplimiento.")
    parser.add_argument("archivo")
    parser.add_argument("--segmento", required=True)
    parser.add_argument("--mes", type=int, required=True)
    parser.add_argument("--real", type=float, required=True)
    args = parser.parse_args()

    ruta = resolver_ruta(args.archivo)
    celdas = leer_celdas(ruta)
    bloques = detectar_bloques(celdas, ETIQUETAS_POR_DEFECTO)
    pares = detectar_columnas(celdas, bloques[0])
    bloque = next((b for b in bloques if b.segmento == args.segmento), None)
    if bloque is None:
        print(f"Segmento no encontrado. Disponibles: {[b.segmento for b in bloques]}")
        return 1
    par_mes = next((p for p in pares if p.tipo_periodo == "mensual" and p.numero_mes == args.mes), None)
    if par_mes is None or par_mes.col_real is None:
        print(f"El mes {args.mes} no existe o no tiene columna de real en el archivo.")
        return 1
    par_acum = next((p for p in pares if p.tipo_periodo == "acumulado" and p.numero_mes == args.mes), None)

    shutil.copy2(ruta, ruta.with_name(ruta.name + ".bak"))
    libro = openpyxl.load_workbook(ruta)
    hoja = libro.worksheets[0]
    escribir_real(hoja, celdas, bloque, par_mes, args.real)
    if par_acum is not None and par_acum.col_real is not None:
        reales = [args.real if p.numero_mes == args.mes else numero_celda(celdas, bloque.filas["meta_driver"], p.col_real)
                  for p in pares if p.tipo_periodo == "mensual" and p.numero_mes <= args.mes and p.col_real]
        if all(r is not None for r in reales):
            escribir_real(hoja, celdas, bloque, par_acum, sum(reales))
    libro.save(ruta)
    print(f"{args.segmento} mes {args.mes}: real = {args.real:g} (respaldo en {ruta.name}.bak)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
