"""Genera un Excel FICTICIO con la misma estructura matricial del archivo de cumplimiento.

Sirve para pruebas automáticas y demostraciones sin usar datos reales.

Uso (desde la raíz del proyecto):
    python scripts/generar_excel_ficticio.py datos/ejemplo/cumplimiento_ficticio.xlsx
"""
from __future__ import annotations

import argparse
import random
from pathlib import Path
from typing import Any

import openpyxl

ETIQUETAS = ("Meta Driver", "Linea Base", "Meta Jugada/Cumplimiento", "Unidades Adicionales")


def generar_datos(segmentos: tuple[str, ...] = ("Segmento Norte", "Segmento Sur"), meses: int = 12,
                  meses_con_real: int = 10, semilla: int = 7) -> dict[str, list[dict[str, Any]]]:
    """Valores mensuales aleatorios pero reproducibles (misma semilla = mismos datos)."""
    rng = random.Random(semilla)
    datos: dict[str, list[dict[str, Any]]] = {}
    for segmento in segmentos:
        filas = []
        for mes in range(1, meses + 1):
            linea_base = rng.randint(800, 1700)
            meta_jugada = rng.randint(linea_base + 50, linea_base + 300)
            filas.append({
                "meta_driver": rng.randint(1500, 3000),
                "linea_base": linea_base,
                "meta_jugada": meta_jugada,
                "real": rng.randint(int(meta_jugada * 0.8), int(meta_jugada * 1.3)) if mes <= meses_con_real else None,
            })
        datos[segmento] = filas
    return datos


def _escribir_par(hoja, col: int, filas: tuple[int, int, int, int], meta_driver: float, linea_base: float,
                  meta_jugada: float, real: float | None, con_columna_real: bool) -> int:
    """Escribe las columnas (meta, real) de un periodo y devuelve la siguiente columna libre."""
    f_md, f_lb, f_mj, f_ua = filas
    hoja.cell(f_md, col, meta_driver)
    hoja.cell(f_lb, col, linea_base)
    hoja.cell(f_mj, col, meta_jugada)
    hoja.cell(f_ua, col, meta_jugada - linea_base)
    if not con_columna_real:
        return col + 1
    real_col = col + 1
    hoja.cell(f_md, real_col, real)
    hoja.merge_cells(start_row=f_md, start_column=real_col, end_row=f_lb, end_column=real_col)
    # Igual que el archivo original: sin real, el % queda en 0 y las unidades en −linea base.
    hoja.cell(f_mj, real_col, (real or 0) / meta_jugada)
    hoja.cell(f_ua, real_col, (real or 0) - linea_base)
    return col + 2


def escribir_libro(ruta: Path, datos: dict[str, list[dict[str, Any]]], columna_inicio: int = 3) -> Path:
    """Escribe el Excel con el patrón: M1, M2, Acum(1-2), M3, Acum(1-3), ..., M(n-1), Acum, Mn."""
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Hoja1"
    fila_base = 1
    for segmento, meses in datos.items():
        filas = (fila_base, fila_base + 1, fila_base + 2, fila_base + 3)
        hoja.cell(fila_base, 1, segmento)
        hoja.merge_cells(start_row=fila_base, start_column=1, end_row=fila_base + 3, end_column=1)
        for desplazamiento, etiqueta in enumerate(ETIQUETAS):
            hoja.cell(fila_base + desplazamiento, 2, etiqueta)

        col = columna_inicio
        acumulado = {"meta_driver": 0, "linea_base": 0, "meta_jugada": 0, "real": 0}
        for mes, d in enumerate(meses, start=1):
            ultimo = mes == len(meses)
            col = _escribir_par(hoja, col, filas, d["meta_driver"], d["linea_base"], d["meta_jugada"],
                                d["real"], con_columna_real=not ultimo)
            for clave in acumulado:
                acumulado[clave] += d[clave] or 0
            if 2 <= mes and not ultimo:
                col = _escribir_par(hoja, col, filas, acumulado["meta_driver"], acumulado["linea_base"],
                                    acumulado["meta_jugada"], acumulado["real"], con_columna_real=True)
        fila_base += 4
    ruta.parent.mkdir(parents=True, exist_ok=True)
    libro.save(ruta)
    return ruta


def main() -> None:
    parser = argparse.ArgumentParser(description="Genera un Excel ficticio de cumplimiento.")
    parser.add_argument("ruta", type=Path, help="Archivo .xlsx a crear")
    parser.add_argument("--semilla", type=int, default=7)
    args = parser.parse_args()
    escribir_libro(args.ruta, generar_datos(semilla=args.semilla))
    print(f"Excel ficticio creado en {args.ruta}")


if __name__ == "__main__":
    main()
