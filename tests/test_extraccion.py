"""Pruebas de la extracción del Excel matricial (con un Excel ficticio de igual estructura)."""
from pathlib import Path

import openpyxl
import pytest

from comun.extractor_base import ErrorExtraccion
from extraccion import (ETIQUETAS_POR_DEFECTO, ExtractorExcelCumplimiento, a_numero, a_slug, detectar_bloques,
                        detectar_columnas, fecha_del_mes, leer_celdas, normalizar_texto)
from generar_excel_ficticio import escribir_libro, generar_datos

RAIZ = Path(__file__).resolve().parents[1]


def extractor(ruta: Path, mes_inicial: int = 1) -> ExtractorExcelCumplimiento:
    return ExtractorExcelCumplimiento(ruta, hoja=None, anio=2026, mes_inicial=mes_inicial)


def test_detecta_bloques_y_columnas(excel_ficticio):
    celdas = leer_celdas(excel_ficticio)
    bloques = detectar_bloques(celdas, ETIQUETAS_POR_DEFECTO)
    assert [b.segmento for b in bloques] == ["Segmento Norte", "Segmento Sur"]
    pares = detectar_columnas(celdas, bloques[0])
    mensuales = [p for p in pares if p.tipo_periodo == "mensual"]
    acumulados = [p for p in pares if p.tipo_periodo == "acumulado"]
    assert [p.numero_mes for p in mensuales] == list(range(1, 13))
    assert [p.numero_mes for p in acumulados] == list(range(2, 12))
    assert mensuales[-1].col_real is None  # el último mes no trae columna de real


def test_omite_meses_sin_real_y_acumulados_incompletos(excel_ficticio):
    observaciones = extractor(excel_ficticio).extraer()
    # Por segmento: 10 meses con real + 9 acumulados completos (2..10), 2 KPI cada uno.
    assert len(observaciones) == 2 * (10 + 9) * 2
    fechas = {o.fecha_referencia for o in observaciones}
    assert "2026-11-01" not in fechas and "2026-12-01" not in fechas


def test_cumplimiento_es_real_sobre_meta_jugada(excel_ficticio):
    datos = generar_datos()
    obs = {o.kpi_id + o.fecha_referencia: o for o in extractor(excel_ficticio).extraer()}
    enero = datos["Segmento Norte"][0]
    cumplimiento = obs["segmento_norte.cumplimiento_mensual2026-01-01"]
    assert cumplimiento.valor == pytest.approx(enero["real"] / enero["meta_jugada"] * 100, abs=1e-3)
    assert cumplimiento.unidad == "porcentaje" and cumplimiento.meta == 100.0
    real = obs["segmento_norte.real_mensual2026-01-01"]
    assert real.valor == enero["real"] and real.meta == enero["meta_jugada"]


def test_acumulado_inconsistente_se_omite(tmp_path):
    ruta = escribir_libro(tmp_path / "x.xlsx", generar_datos())
    libro = openpyxl.load_workbook(ruta)
    hoja = libro.active
    hoja["H1"] = hoja["H1"].value + 999  # real del acumulado meses 1-2 alterado
    libro.save(ruta)
    observaciones = extractor(ruta).extraer()
    claves = {(o.kpi_id, o.fecha_referencia) for o in observaciones}
    assert ("segmento_norte.real_acumulado", "2026-02-01") not in claves
    assert ("segmento_sur.real_acumulado", "2026-02-01") in claves


@pytest.mark.parametrize("anio, mes_inicial, numero_mes, esperado", [
    (2026, 1, 1, "2026-01-01"),
    (2026, 1, 12, "2026-12-01"),
    (2026, 11, 3, "2027-01-01"),
])
def test_fecha_del_mes(anio, mes_inicial, numero_mes, esperado):
    assert fecha_del_mes(anio, mes_inicial, numero_mes) == esperado


def test_etiquetas_con_tildes_y_mayusculas(tmp_path):
    ruta = escribir_libro(tmp_path / "x.xlsx", generar_datos(segmentos=("Único",)))
    libro = openpyxl.load_workbook(ruta)
    libro.active["B2"] = "  LÍNEA   base "
    libro.save(ruta)
    assert detectar_bloques(leer_celdas(ruta), ETIQUETAS_POR_DEFECTO)[0].filas["linea_base"] == 2


def test_falta_una_metrica(tmp_path):
    ruta = escribir_libro(tmp_path / "x.xlsx", generar_datos(segmentos=("A",)))
    libro = openpyxl.load_workbook(ruta)
    libro.active["B4"] = None
    libro.save(ruta)
    with pytest.raises(ErrorExtraccion, match="unidades_adicionales"):
        extractor(ruta).extraer()


def test_archivo_inexistente_o_corrupto(tmp_path):
    with pytest.raises(ErrorExtraccion):
        extractor(tmp_path / "no_existe.xlsx").extraer()
    corrupto = tmp_path / "corrupto.xlsx"
    corrupto.write_text("esto no es un excel")
    with pytest.raises(ErrorExtraccion):
        extractor(corrupto).extraer()


@pytest.mark.parametrize("entrada, esperado", [(5, 5.0), ("1,5", 1.5), (" 2.25 ", 2.25), ("", None),
                                               (None, None), ("abc", None), (True, None), (float("inf"), None)])
def test_a_numero(entrada, esperado):
    assert a_numero(entrada) == esperado


def test_textos():
    assert normalizar_texto("  Meta  JUGADA/Cumplimiento ") == "meta jugada/cumplimiento"
    assert a_slug("Tv + Voz") == "tv_voz"


@pytest.mark.skipif(not (RAIZ / "datos/entrada/Libro1.xlsx").exists(), reason="Libro1.xlsx no disponible")
def test_libro_real_si_existe():
    observaciones = extractor(RAIZ / "datos/entrada/Libro1.xlsx").extraer()
    assert {o.canal for o in observaciones} == {"Pospago", "Tv + Voz"}
    assert all(o.fecha_referencia <= "2026-10-01" for o in observaciones)
